import os
from PyPDF2 import PdfReader
import docx
from langchain.schema import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter

class DocumentProcessor:
    def __init__(self):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200
        )
    
    def load_documents(self, data_path):
        documents = []
        
        for filename in os.listdir(data_path):
            file_path = os.path.join(data_path, filename)
            
            try:
                if filename.endswith('.pdf'):
                    text = self._read_pdf(file_path)
                elif filename.endswith('.docx'):
                    text = self._read_docx(file_path)
                elif filename.endswith('.txt'):
                    text = self._read_txt(file_path)
                else:
                    continue
                
                # Criar documentos LangChain
                docs = self.text_splitter.split_text(text)
                for doc_text in docs:
                    documents.append(Document(
                        page_content=doc_text,
                        metadata={"source": filename}
                    ))
                    
                print(f"✅ Processado: {filename}")
                
            except Exception as e:
                print(f"❌ Erro em {filename}: {e}")
        
        return documents
    
    def _read_pdf(self, file_path):
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text
    
    def _read_docx(self, file_path):
        doc = docx.Document(file_path)
        return "\n".join([paragraph.text for paragraph in doc.paragraphs])
    
    def _read_txt(self, file_path):
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()